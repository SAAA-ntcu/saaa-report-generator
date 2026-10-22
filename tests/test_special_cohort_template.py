"""Unit tests for Special Cohort 7N (E7N & C7N) Chapter 1~3 Templates."""
import unittest
import sys
from pathlib import Path
from docxtpl import DocxTemplate
import docx
from docx.shared import Pt

REPO_ROOT = Path(__file__).resolve().parents[1]
E7N_TPL_PATH = REPO_ROOT / "resources" / "templates" / "special_cohorts" / "E7N_ch1_2_3_template.docx"
C7N_TPL_PATH = REPO_ROOT / "resources" / "templates" / "special_cohorts" / "C7N_ch1_2_3_template.docx"


class TestSpecialCohortTemplate(unittest.TestCase):
    def test_templates_exist(self):
        self.assertTrue(E7N_TPL_PATH.exists(), f"E7N Template not found at {E7N_TPL_PATH}")
        self.assertTrue(C7N_TPL_PATH.exists(), f"C7N Template not found at {C7N_TPL_PATH}")

    def test_e7n_margins_conformance(self):
        doc = docx.Document(str(E7N_TPL_PATH))
        sec = doc.sections[0]
        self.assertAlmostEqual(sec.top_margin.pt, 64.8, delta=0.5)
        self.assertAlmostEqual(sec.bottom_margin.pt, 50.4, delta=0.5)
        self.assertAlmostEqual(sec.left_margin.pt, 56.7, delta=0.5)
        self.assertAlmostEqual(sec.right_margin.pt, 49.6, delta=0.5)

    def test_c7n_margins_conformance(self):
        doc = docx.Document(str(C7N_TPL_PATH))
        sec = doc.sections[0]
        self.assertAlmostEqual(sec.top_margin.pt, 64.8, delta=0.5)
        self.assertAlmostEqual(sec.bottom_margin.pt, 50.4, delta=0.5)
        self.assertAlmostEqual(sec.left_margin.pt, 56.7, delta=0.5)
        self.assertAlmostEqual(sec.right_margin.pt, 49.6, delta=0.5)

    def test_e7n_docxtpl_render_dynamic_loop(self):
        tpl = DocxTemplate(str(E7N_TPL_PATH))
        context = {
            "year": "115",
            "month": "9",
            "day": "25",
            "當年度參與縣市": "桃園市",
            "國字大寫": "一",
            "enrolled_count": "23,273",
            "tested_count": "22,191",
            "boy_count": "11,500",
            "girl_count": "10,691",
            "absence_rate": "4.6",
            "listening_dimensions": "辨識句意、基本問答",
            "listening_dimensions_final": "言談理解",
            "reading_dimensions": "單題",
            "reading_dimensions_final": "題組",
            "framework_rows": [
                {"domain": "聽力", "sub_domain": "辨識句意", "item_numbers": "1、2、3"},
                {"domain": "聽力", "sub_domain": "基本問答", "item_numbers": "4、5、6、7、8"},
                {"domain": "聽力", "sub_domain": "言談理解", "item_numbers": "9、10、11、12、13、14、15"},
                {"domain": "閱讀", "sub_domain": "單題", "item_numbers": "16、17、...、27"},
                {"domain": "閱讀", "sub_domain": "題組", "item_numbers": "28、29、...、35"},
            ],
            "total_items": "35",
        }

        tpl.render(context)
        out_test = REPO_ROOT / "output" / "test_e7n_rendered.docx"
        out_test.parent.mkdir(parents=True, exist_ok=True)
        tpl.save(str(out_test))

        rendered = docx.Document(str(out_test))
        full_text = "\n".join(p.text for p in rendered.paragraphs)
        self.assertIn("學生人數為23,273人，施測人數為22,191人", full_text)
        self.assertIn("男生11,500人、女生10,691人，整體缺考率約為4.6%", full_text)
        self.assertIn("聽力評量向度包含辨識句意、基本問答及言談理解", full_text)
        self.assertIn("閱讀能力則包含單題及題組", full_text)

        table = rendered.tables[0]
        self.assertEqual(len(table.rows), 7)
        self.assertEqual(table.rows[1].cells[0].text.strip(), "聽力")
        self.assertEqual(table.rows[1].cells[1].text.strip(), "辨識句意")
        self.assertEqual(table.rows[1].cells[2].text.strip(), "1、2、3")
        self.assertEqual(table.rows[-1].cells[0].text.strip(), "合計")
        self.assertEqual(table.rows[-1].cells[1].text.strip(), "35題")

    def test_c7n_docxtpl_render_taxonomy_table(self):
        tpl = DocxTemplate(str(C7N_TPL_PATH))
        context = {
            "year": "115",
            "month": "9",
            "day": "25",
            "參與縣市_": "桃園市",
            "參與縣市_final": "",
            "國字大寫": "一",
            "enrolled_count": "23,273",
            "tested_count": "22,223",
            "boy_count": "11,500",
            "girl_count": "10,723",
            "absence_rate": "4.5",
            "knowledge_count": "12",
            "reading_count": "18",
        }

        tpl.render(context)
        out_test = REPO_ROOT / "output" / "test_c7n_rendered.docx"
        out_test.parent.mkdir(parents=True, exist_ok=True)
        tpl.save(str(out_test))

        rendered = docx.Document(str(out_test))
        full_text = "\n".join(p.text for p in rendered.paragraphs)
        self.assertIn("學生人數為23,273人，施測人數為22,223人", full_text)
        self.assertIn("男生11,500人、女生10,723人，整體缺考率約為4.5%", full_text)
        self.assertIn("學力檢測施測之正式題本總題數為30題", full_text)

        table = rendered.tables[0]
        # Total 40 rows: Header + 38 taxonomy rows + 1 total row
        self.assertEqual(len(table.rows), 40)
        total_cell_text = table.rows[39].cells[1].text
        self.assertIn("語文知識12題、閱讀理解18題", total_cell_text)


if __name__ == "__main__":
    unittest.main()
