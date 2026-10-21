import unittest
from pathlib import Path
import tempfile
from docx import Document
import pymupdf

from modules.ch7_appendix import (
    build_chapter_7,
    create_chapter_7_heading,
    create_standalone_chapter_7_docx,
)

class TestCh7Appendix(unittest.TestCase):
    def test_chapter_7_heading_structure(self):
        h = create_chapter_7_heading("柒、附錄")
        self.assertIn("DDEBF7", h.xml)
        self.assertIn("480", h.xml)
        self.assertIn("柒、附錄", h.xml)

    def test_standalone_chapter_7_generation(self):
        # Create a dummy 2-page PDF
        with tempfile.TemporaryDirectory() as tmpdir:
            pdf_path = Path(tmpdir) / "test_booklet.pdf"
            doc_pdf = pymupdf.open()
            for _ in range(2):
                page = doc_pdf.new_page(width=595.3, height=841.9)
                page.draw_rect(pymupdf.Rect(50, 50, 200, 200), color=(0, 0, 0), fill=(0.8, 0.8, 0.8))
            doc_pdf.save(str(pdf_path))
            doc_pdf.close()

            out_docx = Path(tmpdir) / "test_ch7.docx"
            create_standalone_chapter_7_docx(pdf_path, out_docx)

            self.assertTrue(out_docx.exists())
            doc = Document(str(out_docx))
            self.assertEqual(len(doc.paragraphs), 3)  # p0: heading, p1: page 1, p2: page 2
            self.assertEqual(doc.paragraphs[0].text, "柒、附錄")

if __name__ == "__main__":
    unittest.main()
