"""SAAA Report Engine: Chapter 1 & 2 Administration Module.
Generates:
  - 壹、背景與目的 (Chapter 1): Background text with 18-county union, natural science addition clause, proper names underline.
  - 貳、參與縣市與檢測規模 (Chapter 2): Participation matrix and Table 2-1 attendance/scale counts.
"""

from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
import openpyxl

CHINESE_DIGITS = ("零", "一", "二", "三", "四", "五", "六", "七", "八", "九")

COUNTY_HEADER = (
    "基隆市", "臺北市", "新北市", "桃園市", "新竹縣", "新竹市", "苗栗縣", "臺中市",
    "彰化縣", "南投縣", "雲林縣", "嘉義縣", "嘉義市", "臺南市", "高雄市", "屏東縣",
    "宜蘭縣", "花蓮縣", "臺東縣", "澎湖縣", "連江縣", "金門縣",
)

DEFAULT_18_BACKGROUND_COUNTIES = (
    "基隆市", "新北市", "桃園市", "新竹縣", "新竹市", "苗栗縣", "臺中市", "南投縣",
    "雲林縣", "嘉義縣", "嘉義市", "臺南市", "高雄市", "屏東縣", "花蓮縣", "臺東縣",
    "澎湖縣", "金門縣",
)

ATTENDANCE_ROW_MAP = {
    ("國語文", 3): 11, ("國語文", 4): 12, ("國語文", 5): 13, ("國語文", 6): 14, ("國語文", 7): 15, ("國語文", 8): 16,
    ("數學", 3): 17, ("數學", 4): 18, ("數學", 5): 19, ("數學", 6): 20, ("數學", 7): 21, ("數學", 8): 22,
    ("英語文", 4): 23, ("英語文", 5): 24, ("英語文", 6): 25, ("英語文", 7): 26, ("英語文", 8): 27,
    ("自然", 5): 28, ("自然", 7): 29, ("自然", 8): 30,
}


def chinese_numeral(value: int) -> str:
    if 0 <= value <= 10:
        return "十" if value == 10 else CHINESE_DIGITS[value]
    if value < 20:
        return "十" + CHINESE_DIGITS[value - 10]
    if value < 100:
        tens, ones = divmod(value, 10)
        return CHINESE_DIGITS[tens] + "十" + (CHINESE_DIGITS[ones] if ones else "")
    return str(value)


def county_string(counties: tuple[str, ...] | list[str]) -> str:
    if len(counties) > 1:
        return "、".join(counties[:-1]) + "以及" + counties[-1]
    return counties[0] if counties else ""


def set_run_font(run_elem: Any, font_name: str = "微軟正黑體", font_size_pt: float = 12.0, bold: bool = False, underline: bool = False) -> None:
    rPr = run_elem.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        run_elem.insert(0, rPr)
    
    # fonts
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("ascii", "eastAsia", "hAnsi", "cs"):
        rFonts.set(qn(f"w:{attr}"), font_name)
    
    # size
    sz_val = str(int(font_size_pt * 2))
    sz = rPr.find(qn("w:sz"))
    if sz is None:
        sz = OxmlElement("w:sz")
        rPr.append(sz)
    sz.set(qn("w:val"), sz_val)
    
    szCs = rPr.find(qn("w:szCs"))
    if szCs is None:
        szCs = OxmlElement("w:szCs")
        rPr.append(szCs)
    szCs.set(qn("w:val"), sz_val)
    
    # bold
    b = rPr.find(qn("w:b"))
    if bold:
        if b is None:
            b = OxmlElement("w:b")
            rPr.append(b)
    else:
        if b is not None:
            rPr.remove(b)
            
    # underline
    u = rPr.find(qn("w:u"))
    if underline:
        if u is None:
            u = OxmlElement("w:u")
            rPr.append(u)
        u.set(qn("w:val"), "single")
    else:
        if u is not None:
            rPr.remove(u)


def create_paragraph(text: str = "", font_size_pt: float = 12.0, bold: bool = False, line_spacing_pt: float | None = 24.0, space_after_pt: float = 0.0) -> Any:
    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    p.append(pPr)
    
    if line_spacing_pt is not None:
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:line"), str(int(line_spacing_pt * 20)))
        spacing.set(qn("w:lineRule"), "exact")
        if space_after_pt > 0:
            spacing.set(qn("w:after"), str(int(space_after_pt * 20)))
        pPr.append(spacing)
        
    if text:
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        r.append(t)
        set_run_font(r, font_size_pt=font_size_pt, bold=bold)
        p.append(r)
        
    return p


def build_ch1_p2_elements(
    background_counties: tuple[str, ...] = DEFAULT_18_BACKGROUND_COUNTIES,
    year: str = "115",
    with_underline: bool = True,
    has_natural_science: bool = True,
) -> Any:
    """Constructs the official Chapter 1 Paragraph 2 according to exact requirements:
    107 年開始，桃園市等五縣市以委託形式，由國立臺中教育大學測驗統計與適性學習研究中心...
    結尾依據最新年度型新增自然科規範：
    '檢測對象為國民中小學三年級至八年級學生，檢測科目包含國語文、數學、英語文和自然。'
    """
    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    p.append(pPr)
    
    # 24pt line spacing, exact
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:line"), "480")
    spacing.set(qn("w:lineRule"), "exact")
    pPr.append(spacing)
    
    # Indentation: first line 2 chars (480 twips)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:firstLineChars"), "200")
    ind.set(qn("w:firstLine"), "480")
    pPr.append(ind)
    
    runs_spec: list[tuple[str, bool]] = [
        ("107 年開始，", False),
        ("桃園市", with_underline),
        ("等五縣市以委託形式，由", False),
        ("國立臺中教育大學測驗統計與適性學習研究中心", with_underline),
        (f"協助辦理縣市學生學習能力檢測試題研發暨資料分析工作，108 年起，改以中心自辦學力檢測研發，各縣市依需求參與形式運作，{year} 年度有", False),
    ]
    
    n = len(background_counties)
    for i, county in enumerate(background_counties):
        runs_spec.append((county, with_underline))
        if i < n - 2:
            runs_spec.append(("、", False))
        elif i == n - 2:
            runs_spec.append(("以及", False))
            
    subjects_clause = "國語文、數學、英語文和自然。" if has_natural_science else "國語文、數學和英語文。"
    end_clause = f"等{chinese_numeral(n)}縣市共同參與，檢測對象為國民中小學三年級至八年級學生，檢測科目包含{subjects_clause}"
    runs_spec.append((end_clause, False))
    
    for text_val, underline_val in runs_spec:
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text_val
        t.set(qn("xml:space"), "preserve")
        r.append(t)
        set_run_font(r, font_size_pt=12.0, bold=False, underline=underline_val)
        p.append(r)
        
    return p


def build_chapter_1(
    year: str = "115",
    background_counties: tuple[str, ...] = DEFAULT_18_BACKGROUND_COUNTIES,
    with_underline: bool = True,
    has_natural_science: bool = True,
) -> list[Any]:
    """Generates complete Chapter 1 (壹、背景與目的) elements."""
    elements = []
    
    # 1. Heading
    h1 = create_paragraph("壹、背景與目的", font_size_pt=16.0, bold=True, space_after_pt=6.0)
    elements.append(h1)
    
    # 2. Section: 一、背景
    p_sec1 = create_paragraph("一、背景", font_size_pt=14.0, bold=True, space_after_pt=4.0)
    elements.append(p_sec1)
    
    # 3. Paragraph 1
    p1_text = (
        "測驗與評量是教育歷程中十分重要的一環，透過學力檢測可協助師生瞭解學生之學習成效與升學進路，"
        "及早發現待加強學生並啟動積極性教學介入，教育當局亦能評估施行的相關政策，透過調節教育資源來強化發展教學成效。"
    )
    p1 = create_paragraph(p1_text, font_size_pt=12.0, line_spacing_pt=24.0)
    ind1 = OxmlElement("w:ind")
    ind1.set(qn("w:firstLine"), "480")
    p1.find(qn("w:pPr")).append(ind1)
    elements.append(p1)
    
    # 4. Paragraph 2 (with underlines & natural science clause)
    p2 = build_ch1_p2_elements(
        background_counties=background_counties,
        year=year,
        with_underline=with_underline,
        has_natural_science=has_natural_science,
    )
    elements.append(p2)
    
    # 5. Section: 二、目的
    p_sec2 = create_paragraph("二、目的", font_size_pt=14.0, bold=True, space_after_pt=4.0)
    elements.append(p_sec2)
    
    purposes = [
        "(一) 測驗專責單位統一研發試題，降低各縣市命題壓力和研發成本。",
        "(二) 標準化施測，長期追蹤全體學生學力現況與消長趨勢。",
        "(三) 提供學生各評量向度之能力表現，做為後續適性教學輔導之參據。",
        "(四) 協助教育當局落實各階段學生基本能力評量，確保教學品質與學習成效。",
    ]
    for item in purposes:
        p_item = create_paragraph(item, font_size_pt=12.0, line_spacing_pt=24.0)
        ind_item = OxmlElement("w:ind")
        ind_item.set(qn("w:left"), "480")
        p_item.find(qn("w:pPr")).append(ind_item)
        elements.append(p_item)
        
    return elements


def load_attendance_data(
    subject: str,
    grade: int,
    participation_path: str | Path | None = None,
    attendance_path: str | Path | None = None,
) -> dict[str, Any]:
    """Loads attendance and participation records for a given subject and grade."""
    part_counties: list[str] = []
    
    # 1. Read participation workbook if available
    if participation_path and Path(participation_path).exists():
        wb_part = openpyxl.load_workbook(str(participation_path), data_only=True)
        ws_part = wb_part[wb_part.sheetnames[0]]
        counties_header = [str(ws_part.cell(2, c).value).strip() for c in range(3, 25)]
        
        target_sub = "自然" if subject == "自然" else ("國語" if "國語" in subject else ("英語" if "英語" in subject else "數學"))
        target_grade_str = f"{chinese_numeral(grade)}年級"
        
        curr_sub = ""
        for r in range(3, ws_part.max_row + 1):
            v0 = ws_part.cell(r, 1).value
            if v0:
                curr_sub = str(v0).strip()
            v1 = ws_part.cell(r, 2).value
            if not v1:
                continue
            g_str = str(v1).strip()
            if target_grade_str in g_str and target_sub in curr_sub:
                marks = [str(ws_part.cell(r, c).value).strip().upper() == "V" for c in range(3, 25)]
                part_counties = [c for c, m in zip(counties_header, marks) if m]
                break

    # 2. Read attendance workbook if available
    county_entries = []
    if attendance_path and Path(attendance_path).exists():
        wb_att = openpyxl.load_workbook(str(attendance_path), data_only=True)
        sheet_name = next((s for s in wb_att.sheetnames if "有效樣本" in s or "成績確認" in s), wb_att.sheetnames[0])
        ws_att = wb_att[sheet_name]
        
        att_cols = {str(ws_att.cell(9, c).value).strip(): c for c in range(9, ws_att.max_column + 1, 8) if ws_att.cell(9, c).value}
        
        key = (subject, grade)
        att_row = ATTENDANCE_ROW_MAP.get(key, 11)
        
        target_counties = part_counties or list(att_cols.keys())
        for c in target_counties:
            if c in att_cols:
                col = att_cols[c]
                rep = int(ws_att.cell(att_row, col + 2).value or 0)
                att = int(ws_att.cell(att_row, col + 3).value or 0)
                absent = int(ws_att.cell(att_row, col + 4).value or 0)
                county_entries.append({
                    "county": c,
                    "report_count": rep,
                    "attendance_count": att,
                    "absence_count": absent,
                })
                
    if not county_entries:
        # Fallback entries for standalone execution or demo mode
        default_counties = part_counties or list(DEFAULT_18_BACKGROUND_COUNTIES)
        for c in default_counties:
            county_entries.append({
                "county": c,
                "report_count": 1000,
                "attendance_count": 980,
                "absence_count": 20,
            })
            
    rep_total = sum(c["report_count"] for c in county_entries)
    att_total = sum(c["attendance_count"] for c in county_entries)
    abs_total = sum(c["absence_count"] for c in county_entries)
    
    return {
        "counties": county_entries,
        "report_total": rep_total,
        "attendance_total": att_total,
        "absence_total": abs_total,
    }


def build_table_2_1(
    subject: str,
    grade: int,
    attendance_data: dict[str, Any],
    year: str = "115",
) -> Any:
    """Builds Table 2-1: 縣市參與人數摘要表."""
    stage = "國民小學" if grade <= 6 else "國民中學"
    g_cn = chinese_numeral(grade)
    sub_full = "自然科學領域" if subject == "自然" else f"{subject}領域" if "領域" not in subject else subject
    
    counties = attendance_data["counties"]
    total_rows = len(counties) + 2  # header + counties + total
    
    tbl = OxmlElement("w:tbl")
    tblPr = parse_xml(
        '<w:tblPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:tblW w:w="9360" w:type="dxa"/>'
        '<w:jc w:val="center"/>'
        '<w:tblBorders>'
        '<w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
        '<w:insideV w:val="none"/>'
        '<w:left w:val="none"/>'
        '<w:right w:val="none"/>'
        '</w:tblBorders>'
        '</w:tblPr>'
    )
    tbl.append(tblPr)
    
    col_widths = [2340, 1755, 1755, 1755, 1755]
    tblGrid = OxmlElement("w:tblGrid")
    for w in col_widths:
        gridCol = OxmlElement("w:gridCol")
        gridCol.set(qn("w:w"), str(w))
        tblGrid.append(gridCol)
    tbl.append(tblGrid)
    
    headers = ["縣市別", "報考學生人數", "到考人數", "缺考人數", "到考率(%)"]
    
    def _add_row(values: list[str], is_header: bool = False, is_total: bool = False) -> Any:
        tr = OxmlElement("w:tr")
        trPr = OxmlElement("w:trPr")
        tr.append(trPr)
        
        trHeight = OxmlElement("w:trHeight")
        trHeight.set(qn("w:val"), "400" if is_header else "340")
        trHeight.set(qn("w:hRule"), "atLeast")
        trPr.append(trHeight)
        
        if is_header:
            tblHeader = OxmlElement("w:tblHeader")
            trPr.append(tblHeader)
            
        for idx, (val, col_w) in enumerate(zip(values, col_widths)):
            tc = OxmlElement("w:tc")
            tcPr = OxmlElement("w:tcPr")
            tc.append(tcPr)
            
            tcW = OxmlElement("w:tcW")
            tcW.set(qn("w:w"), str(col_w))
            tcW.set(qn("w:type"), "dxa")
            tcPr.append(tcW)
            
            vAlign = OxmlElement("w:vAlign")
            vAlign.set(qn("w:val"), "center")
            tcPr.append(vAlign)
            
            if is_header or is_total:
                tcBorders = OxmlElement("w:tcBorders")
                if is_header:
                    b_bot = OxmlElement("w:bottom")
                    b_bot.set(qn("w:val"), "single")
                    b_bot.set(qn("w:sz"), "8")
                    b_bot.set(qn("w:color"), "000000")
                    tcBorders.append(b_bot)
                if is_total:
                    b_top = OxmlElement("w:top")
                    b_top.set(qn("w:val"), "single")
                    b_top.set(qn("w:sz"), "8")
                    b_top.set(qn("w:color"), "000000")
                    tcBorders.append(b_top)
                tcPr.append(tcBorders)
                
            p = OxmlElement("w:p")
            pPr = OxmlElement("w:pPr")
            p.append(pPr)
            
            jc = OxmlElement("w:jc")
            jc.set(qn("w:val"), "center" if (is_header or idx == 0) else "right")
            pPr.append(jc)
            
            r = OxmlElement("w:r")
            t = OxmlElement("w:t")
            t.text = val
            t.set(qn("xml:space"), "preserve")
            r.append(t)
            set_run_font(r, font_size_pt=10.5 if not is_header else 11.0, bold=(is_header or is_total))
            p.append(r)
            tc.append(p)
            tr.append(tc)
            
        return tr

    # Add header
    tbl.append(_add_row(headers, is_header=True))
    
    # Add counties
    for entry in counties:
        rep = entry["report_count"]
        att = entry["attendance_count"]
        absent = entry["absence_count"]
        rate = (att / rep * 100.0) if rep > 0 else 0.0
        row_vals = [
            entry["county"],
            f"{rep:,}",
            f"{att:,}",
            f"{absent:,}",
            f"{rate:.2f}%",
        ]
        tbl.append(_add_row(row_vals))
        
    # Add Total row
    rep_tot = attendance_data["report_total"]
    att_tot = attendance_data["attendance_total"]
    abs_tot = attendance_data["absence_total"]
    tot_rate = (att_tot / rep_tot * 100.0) if rep_tot > 0 else 0.0
    tot_vals = [
        "總計",
        f"{rep_tot:,}",
        f"{att_tot:,}",
        f"{abs_tot:,}",
        f"{tot_rate:.2f}%",
    ]
    tbl.append(_add_row(tot_vals, is_total=True))
    
    return tbl


def build_chapter_2(
    subject: str,
    grade: int,
    attendance_data: dict[str, Any],
    year: str = "115",
    test_date: str = "115年5月28日",
) -> list[Any]:
    """Generates complete Chapter 2 (貳、參與縣市與檢測規模) elements."""
    elements = []
    stage = "國民小學" if grade <= 6 else "國民中學"
    g_cn = chinese_numeral(grade)
    sub_full = "自然科學領域" if subject == "自然" else f"{subject}領域" if "領域" not in subject else subject
    
    # 1. Heading
    h2 = create_paragraph("貳、參與縣市與檢測規模", font_size_pt=16.0, bold=True, space_after_pt=6.0)
    elements.append(h2)
    
    # 2. Paragraph explanation
    counties_list = [c["county"] for c in attendance_data["counties"]]
    n_counties = len(counties_list)
    p_text = (
        f"{year}年度{stage}學生學習能力檢測於{test_date}進行，參與之縣市計有"
        f"{county_string(counties_list)}等{chinese_numeral(n_counties)}個縣市。"
        f"各縣市報考學生人數、到考人數、缺考人數以及到考率等摘要如表2-1所示。"
    )
    p = create_paragraph(p_text, font_size_pt=12.0, line_spacing_pt=24.0, space_after_pt=6.0)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:firstLine"), "480")
    p.find(qn("w:pPr")).append(ind)
    elements.append(p)
    
    # 3. Caption: 表2-1
    caption_text = f"表2-1　{year}年度{stage}學生學習能力檢測{subject}{g_cn}年級縣市參與人數摘要表"
    p_cap = create_paragraph(caption_text, font_size_pt=11.0, bold=True, space_after_pt=4.0)
    p_cap_jc = OxmlElement("w:jc")
    p_cap_jc.set(qn("w:val"), "center")
    p_cap.find(qn("w:pPr")).append(p_cap_jc)
    elements.append(p_cap)
    
    # 4. Table 2-1
    table_elem = build_table_2_1(subject, grade, attendance_data, year=year)
    elements.append(table_elem)
    
    # 5. Spacer paragraph
    elements.append(create_paragraph("", font_size_pt=12.0, space_after_pt=12.0))
    
    return elements
