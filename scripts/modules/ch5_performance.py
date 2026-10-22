"""SAAA Report Engine: Chapter 5 Performance Analysis Module.
Generates:
  - 伍、整體學力表現分析 (Chapter 5)
  - Section 1: 不同背景變項學生整體學力表現分析 (Figure 5-1 distribution chart)
  - Section 2: 學生於不同學習重點之答對率表現 (Table 5-1-1 ~ 5-1-N cross-tabulations)
  - Live 108 Curriculum PDF descriptor extraction using PyMuPDF.
"""

from __future__ import annotations

from dataclasses import dataclass
from functools import cache
from pathlib import Path
import re
from typing import Any

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
import openpyxl
import pymupdf

from .ch1_2_admin import chinese_numeral, create_paragraph, set_run_font


INTERVALS = ("91~100", "81~90", "71~80", "61~70", "51~60", "41~50", "31~40", "21~30", "11~20", "0~10")

CANONICAL_AREAS: dict[str, list[str]] = {
    "國語文": [
        "形音知識", "字詞知識", "語法知識", "修辭知識", "章法知識", "文體知識",
        "國學知識", "字詞理解", "句子理解", "段落理解", "篇章理解",
    ],
    "數學_小學": [
        "數與計算", "量與實測", "空間與形狀", "資料與不確定性", "關係",
    ],
    "數學_中學": [
        "數與量", "空間與形狀與坐標幾何", "代數與函數", "資料與不確定性",
    ],
    "英語文_小學": [
        "聽力-語音聽辨", "聽力-辭彙聽辨", "聽力-教室生活辭句理解", "聽力-教室生活辭句理解與回應", "聽力-文化節慶理解",
        "閱讀能力-字詞辨識", "閱讀能力-句子理解", "閱讀能力-短文與對話理解", "閱讀能力-文化節慶理解",
    ],
    "英語文_中學": [
        "聽力-常用單句", "聽力-日常對答", "聽力-信息推論", "聽力-生活對話與日常信息",
        "閱讀能力-情境字彙", "閱讀能力-情境文法與句型", "閱讀能力-短文推論及圖文轉換", "閱讀能力-綜合能力",
        "閱讀能力-理解篇章主旨或細節", "閱讀能力-非連續文本理解或推論",
    ],
    "自然_5": [
        "交互作用", "改變與穩定", "構造與功能", "物質與能量", "科學與生活", "系統與尺度", "資源與永續性",
    ],
    "自然_7": [
        "能量的形態與流動", "生物的構造與功能", "地球環境", "演化與延續", "生物與環境", "跨科主題",
    ],
    "自然_8": [
        "物質的組成與特性", "能量的形式、轉換及流動", "物質的結構與功能", "物質系統", "物質的反應、平衡及製造", "自然界的現象與交互作用",
    ],
}


@dataclass(frozen=True, slots=True)
class Item:
    number: int
    rate: float
    codes: tuple[str, ...]
    area: str = ""


def clean_description(value: str) -> str:
    description = " ".join(value.split())
    for width in range(min(80, len(description) // 2), 9, -1):
        found_duplicate = False
        for start in range(len(description) - width):
            fragment = description[start:start + width]
            second = description.find(fragment, start + width)
            if second >= 0:
                description = description[:second] + description[second + width:]
                found_duplicate = True
                break
        if found_duplicate:
            break
    return description


@cache
def extract_curriculum_descriptions(curriculum_pdf_path: str | Path, subject: str) -> dict[str, str]:
    """Extracts curriculum code descriptions from official Taiwan 108 Curriculum PDFs."""
    pdf_path = Path(curriculum_pdf_path)
    if not pdf_path.exists():
        return {}
        
    descriptions: dict[str, str] = {}
    
    if "數學" in subject or subject == "M":
        pattern = re.compile(r"(?<![A-Za-z0-9])([A-Z]-[3-8]-\d+)(?![A-Za-z0-9])")
        with pymupdf.open(str(pdf_path)) as document:
            for page in document:
                text = page.get_text()
                if not pattern.search(text):
                    continue
                # Try table extraction first
                tables = page.find_tables().tables
                for table in tables:
                    for row in table.extract():
                        if len(row) < 2:
                            continue
                        matches = tuple(pattern.finditer(row[0] or ""))
                        for idx, m in enumerate(matches):
                            end = matches[idx + 1].start() if idx + 1 < len(matches) else len(row[0] or "")
                            if end > m.end():
                                continue
                            desc = clean_description(row[1] or "")
                            if desc:
                                descriptions.setdefault(m.group(1), desc[:280])
                                
        if not descriptions:
            # fallback regex on plain text
            with pymupdf.open(str(pdf_path)) as document:
                text = "\n".join(p.get_text() for p in document)
                matches = tuple(pattern.finditer(text))
                for idx, m in enumerate(matches):
                    end = matches[idx + 1].start() if idx + 1 < len(matches) else min(len(text), m.end() + 240)
                    desc = clean_description(text[m.end():end])
                    if desc:
                        descriptions.setdefault(m.group(1), desc[:280])
        return descriptions

    if "自然" in subject or subject == "S":
        pattern = re.compile(r"(?<![A-Za-z0-9])((?:IN[a-z]|[A-Z][a-z]?)-[ⅠⅡⅢⅣⅤ]-\d+)(?![A-Za-z0-9])")
        with pymupdf.open(str(pdf_path)) as document:
            for page in document:
                if not pattern.search(page.get_text()):
                    continue
                for table in page.find_tables().tables:
                    for row in table.extract():
                        for cell in row:
                            matches = tuple(pattern.finditer(cell or ""))
                            for idx, m in enumerate(matches):
                                end = matches[idx + 1].start() if idx + 1 < len(matches) else len(cell or "")
                                desc = clean_description((cell or "")[m.end():end])
                                if desc:
                                    descriptions.setdefault(m.group(1), desc[:280])
        descriptions["Mc-Ⅳ-0"] = "科學在生活中的應用（Mc；來源學習內容表未提供有效序號）"
        return descriptions

    if "英語" in subject or subject == "E":
        pattern = re.compile(r"(?<![A-Za-z0-9])([1-9]-[ⅠⅡⅢⅣⅤ]-\d+)(?![A-Za-z0-9])")
        with pymupdf.open(str(pdf_path)) as document:
            text = "\n".join(page.get_text() for page in document)
            matches = tuple(pattern.finditer(text))
            for idx, m in enumerate(matches):
                end = matches[idx + 1].start() if idx + 1 < len(matches) else min(len(text), m.end() + 240)
                desc = clean_description(text[m.end():end])
                desc = re.sub(r"^(?:學習內容|學習表現|編碼)+\s*", "", desc)
                if desc:
                    descriptions.setdefault(m.group(1), desc[:280])
        return descriptions

    # Chinese
    pattern = re.compile(r"(?<![A-Za-z0-9])([1-6]-[ⅠⅡⅢⅣⅤ]-\d+|[A-Za-z]{1,2}-[ⅠⅡⅢⅣⅤ]-\d+)(?![A-Za-z0-9])")
    with pymupdf.open(str(pdf_path)) as document:
        text = "\n".join(page.get_text() for page in document)
        matches = tuple(pattern.finditer(text))
        for idx, m in enumerate(matches):
            end = matches[idx + 1].start() if idx + 1 < len(matches) else min(len(text), m.end() + 240)
            desc = clean_description(text[m.end():end])
            if desc:
                descriptions.setdefault(m.group(1), desc[:280])
    return descriptions


def _make_cell(
    width: int,
    paragraphs: list[BaseOxmlElement],
    *,
    shd_fill: str | None = None,
    v_merge: str | None = None,
    grid_span: int | None = None,
) -> BaseOxmlElement:
    cell = OxmlElement("w:tc")
    cell_props = OxmlElement("w:tcPr")
    
    cell_width = OxmlElement("w:tcW")
    cell_width.set(qn("w:w"), str(width))
    cell_width.set(qn("w:type"), "dxa")
    cell_props.append(cell_width)
    
    if grid_span and grid_span > 1:
        gs = OxmlElement("w:gridSpan")
        gs.set(qn("w:val"), str(grid_span))
        cell_props.append(gs)

    if shd_fill:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), shd_fill)
        cell_props.append(shd)
        
    v_align = OxmlElement("w:vAlign")
    v_align.set(qn("w:val"), "center")
    cell_props.append(v_align)
    
    if v_merge:
        vm = OxmlElement("w:vMerge")
        if v_merge != "continue":
            vm.set(qn("w:val"), v_merge)
        cell_props.append(vm)
        
    cell.append(cell_props)
    for p in paragraphs:
        cell.append(p)
    return cell


def _cell_paragraph(text: str, *, bold: bool = False, center: bool = False) -> BaseOxmlElement:
    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:line"), "440")
    spacing.set(qn("w:lineRule"), "exact")
    pPr.append(spacing)
    if center:
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "center")
        pPr.append(jc)
    p.append(pPr)
    if text:
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        r.append(t)
        set_run_font(r, font_size_pt=10.5, bold=bold)
        p.append(r)
    return p


def build_area_cross_table(
    items: list[Item],
    descriptions: dict[str, str],
    is_chinese: bool = False,
) -> BaseOxmlElement:
    """Builds a Table 5-1-x 10-interval pass rate cross-tabulation table.
    - If is_chinese=True: 3-column table (答對率%, 題號, 學習重點: 代碼\\t說明)
    - If is_chinese=False: 4-column table (答對率%, 題號, 代碼, 說明; header columns 2&3 merged)
    """
    if is_chinese:
        widths = (1257, 752, 7772)
    else:
        # Col 0: 答對率% (1139), Col 1: 題號 (709), Col 2: 代碼 (992), Col 3: 說明 (6941)
        widths = (1139, 709, 992, 6941)
        
    table = OxmlElement("w:tbl")
    props = OxmlElement("w:tblPr")
    
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), str(sum(widths)))
    width.set(qn("w:type"), "dxa")
    props.append(width)
    
    tbl_jc = OxmlElement("w:jc")
    tbl_jc.set(qn("w:val"), "center")
    props.append(tbl_jc)
    
    borders = OxmlElement("w:tblBorders")
    for edge, size in (("top", "12"), ("bottom", "12"), ("insideH", "4")):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:color"), "BFBFBF")
        borders.append(border)
    for edge in ("left", "right", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "none")
        borders.append(border)
    props.append(borders)
    table.append(props)
    
    grid = OxmlElement("w:tblGrid")
    for val in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(val))
        grid.append(col)
    table.append(grid)
    
    # Header row
    hdr_row = OxmlElement("w:tr")
    hdr_pr = OxmlElement("w:trPr")
    hdr_pr.append(OxmlElement("w:tblHeader"))
    hdr_pr.append(OxmlElement("w:cantSplit"))
    tr_h = OxmlElement("w:trHeight")
    tr_h.set(qn("w:val"), "400")
    hdr_pr.append(tr_h)
    hdr_row.append(hdr_pr)
    
    if is_chinese:
        hdr_row.append(_make_cell(widths[0], [_cell_paragraph("答對率%", bold=True, center=True)]))
        hdr_row.append(_make_cell(widths[1], [_cell_paragraph("題號", bold=True, center=True)]))
        hdr_row.append(_make_cell(widths[2], [_cell_paragraph("學習重點", bold=True, center=False)]))
    else:
        hdr_row.append(_make_cell(widths[0], [_cell_paragraph("答對率%", bold=True, center=True)]))
        hdr_row.append(_make_cell(widths[1], [_cell_paragraph("題號", bold=True, center=True)]))
        # Merge header cells 2 & 3 for Math / English / Science
        merged_w = widths[2] + widths[3]
        hdr_row.append(_make_cell(merged_w, [_cell_paragraph("學習重點", bold=True, center=False)], grid_span=2))
    table.append(hdr_row)
    
    # Group items into intervals
    by_interval: dict[str, list[Item]] = {label: [] for label in INTERVALS}
    for item in items:
        percent = round(item.rate * 100) if item.rate <= 1.0 else round(item.rate)
        iv = next((lbl for ceil_val, lbl in zip(range(90, 0, -10), INTERVALS) if percent > ceil_val), INTERVALS[-1])
        by_interval[iv].append(item)
        
    for iv_idx, iv in enumerate(INTERVALS):
        shd_fill = "DDEBF7" if iv_idx % 2 == 0 else "auto"
        interval_items = by_interval[iv]
        # Sort items in interval by question number
        interval_items.sort(key=lambda it: it.number)
        
        if not interval_items:
            row = OxmlElement("w:tr")
            tr_pr = OxmlElement("w:trPr")
            tr_pr.append(OxmlElement("w:cantSplit"))
            row.append(tr_pr)
            if is_chinese:
                row.append(_make_cell(widths[0], [_cell_paragraph(iv, center=True)], shd_fill=shd_fill))
                row.append(_make_cell(widths[1], [_cell_paragraph("", center=True)], shd_fill=shd_fill))
                row.append(_make_cell(widths[2], [_cell_paragraph("")], shd_fill=shd_fill))
            else:
                row.append(_make_cell(widths[0], [_cell_paragraph(iv, center=True)], shd_fill=shd_fill))
                row.append(_make_cell(widths[1], [_cell_paragraph("", center=True)], shd_fill=shd_fill))
                row.append(_make_cell(widths[2], [_cell_paragraph("", center=True)], shd_fill=shd_fill))
                row.append(_make_cell(widths[3], [_cell_paragraph("")], shd_fill=shd_fill))
            table.append(row)
        else:
            for item_idx, itm in enumerate(interval_items):
                row = OxmlElement("w:tr")
                tr_pr = OxmlElement("w:trPr")
                tr_pr.append(OxmlElement("w:cantSplit"))
                row.append(tr_pr)
                
                # First col: interval (vertically merged)
                v_merge = "restart" if item_idx == 0 else "continue"
                row.append(_make_cell(
                    widths[0],
                    [_cell_paragraph(iv if item_idx == 0 else "", center=True)],
                    shd_fill=shd_fill,
                    v_merge=v_merge,
                ))
                
                # Second col: question number
                row.append(_make_cell(
                    widths[1],
                    [_cell_paragraph(str(itm.number), center=True)],
                    shd_fill=shd_fill,
                ))
                
                if is_chinese:
                    # Third col: code + Tab + description
                    desc_runs = []
                    for code in itm.codes:
                        clean_code = code.strip()
                        desc = descriptions.get(clean_code, "")
                        if desc:
                            desc_runs.append(f"{clean_code}\t{desc}")
                        else:
                            desc_runs.append(clean_code)
                    full_desc_text = "\n".join(desc_runs)
                    row.append(_make_cell(
                        widths[2],
                        [_cell_paragraph(full_desc_text)],
                        shd_fill=shd_fill,
                    ))
                else:
                    # Third col: code only
                    codes_str = "\n".join(c.strip() for c in itm.codes if c.strip())
                    row.append(_make_cell(
                        widths[2],
                        [_cell_paragraph(codes_str, center=True)],
                        shd_fill=shd_fill,
                    ))
                    # Fourth col: description only
                    desc_runs = []
                    for code in itm.codes:
                        clean_code = code.strip()
                        desc = descriptions.get(clean_code, "")
                        if desc:
                            desc_runs.append(desc)
                    full_desc_text = "\n".join(desc_runs)
                    row.append(_make_cell(
                        widths[3],
                        [_cell_paragraph(full_desc_text)],
                        shd_fill=shd_fill,
                    ))
                table.append(row)
                
    return table


def sort_areas(area_names: list[str], subject: str, grade_num: int, is_7n: bool = False) -> list[str]:
    """Sorts content areas according to official 108 Curriculum sequence."""
    key = ""
    if "國語" in subject or subject == "C":
        key = "國語文"
    elif "數學" in subject or subject == "M":
        key = "數學_小學" if (grade_num <= 6 or is_7n) else "數學_中學"
    elif "英語" in subject or subject == "E":
        key = "英語文_小學" if (grade_num <= 6 and not is_7n) else "英語文_中學"
    elif "自然" in subject or subject == "S":
        key = f"自然_{grade_num}" if grade_num in (5, 7, 8) else "自然_5"
        
    canonical = CANONICAL_AREAS.get(key, [])
    
    def sort_key(name: str):
        for idx, c_name in enumerate(canonical):
            if c_name == name or c_name in name or name in c_name:
                return (0, idx)
        return (1, name)
        
    return sorted(area_names, key=sort_key)


def build_chapter_5(
    subject: str,
    grade: int | str,
    items: list[Item],
    descriptions: dict[str, str],
    figure_path: str | Path | None = None,
    year: str = "115",
    is_7n: bool = False,
) -> list[Any]:
    """Generates complete Chapter 5 (伍、整體學力表現分析) elements."""
    elements = []
    is_7n = is_7n or "7N" in str(grade).upper() or "新生" in str(grade)
    grade_num = int(re.sub(r"\D", "", str(grade)) or 7)
    g_cn = chinese_numeral(grade_num)
    stage = "國民中學" if (grade_num >= 7 or is_7n) else "國民小學"
    cohort_suffix = "新生" if is_7n else "學生"
    city_prefix = "桃園市" if is_7n else ""
    is_chinese = "國語" in subject or subject == "C"
    
    # 1. Heading: 伍、整體學力表現分析
    h5 = create_paragraph("伍、整體學力表現分析", font_size_pt=16.0, bold=True, space_after_pt=6.0)
    elements.append(h5)
    
    # 2. Section 1: 一、不同背景變項學生整體學力表現分析
    p_sec1 = create_paragraph("一、不同背景變項學生整體學力表現分析", font_size_pt=14.0, bold=True, space_after_pt=4.0)
    elements.append(p_sec1)
    
    cohort_tag = f"{g_cn}年級{cohort_suffix}"
    fig_tag = f"{g_cn}年級新生" if is_7n else f"{g_cn}年級"
    sec1_text = (
        f"本計畫同時採用古典測驗理論（classical test theory, CTT）與試題反應理論（item response theory, IRT）"
        f"三參數模式作為試題分析理論基礎，並以三參數模式進行學生能力估計。三參數模式有a、b、c三種參數，"
        f"分別代表鑑別度、難度以及猜測度。學生能力估計值則參考PISA與TIMSS等國際大型測驗的作法，"
        f"將估計出來的能力轉換至平均分數500，標準差100的量尺上，再以轉換所得量尺分數進行不同背景變項學生表現分析。"
        f"本年度{subject}{cohort_tag}整體學力表現分布狀況，如圖5-1所示。"
    )
    p_desc1 = create_paragraph(sec1_text, font_size_pt=12.0, line_spacing_pt=24.0, space_after_pt=6.0)
    ind1 = OxmlElement("w:ind")
    ind1.set(qn("w:firstLine"), "480")
    p_desc1.find(qn("w:pPr")).append(ind1)
    elements.append(p_desc1)
    
    # Caption: 圖5-1
    fig_caption = f"圖5-1　{subject}{fig_tag}整體學生學力表現分布圖"
    p_fig_cap = create_paragraph(fig_caption, font_size_pt=11.0, bold=True, space_after_pt=4.0)
    jc_fig = OxmlElement("w:jc")
    jc_fig.set(qn("w:val"), "center")
    p_fig_cap.find(qn("w:pPr")).append(jc_fig)
    elements.append(p_fig_cap)
    
    # Insert Figure image if provided and exists
    if figure_path and Path(figure_path).exists():
        temp_doc = Document()
        img_p = temp_doc.add_paragraph()
        img_p.alignment = 1  # center
        run = img_p.add_run()
        run.add_picture(str(figure_path))
        elements.append(img_p._p)
    else:
        p_ph = create_paragraph("［學力表現分布圖檔待插入］", font_size_pt=11.0, space_after_pt=6.0)
        jc_ph = OxmlElement("w:jc")
        jc_ph.set(qn("w:val"), "center")
        p_ph.find(qn("w:pPr")).append(jc_ph)
        elements.append(p_ph)
        
    # 3. Section 2: 二、學生於不同學習重點之答對率表現
    p_sec2 = create_paragraph(f"二、{stage}{cohort_tag}於不同學習重點之答對率表現", font_size_pt=14.0, bold=True, space_after_pt=4.0)
    elements.append(p_sec2)
    
    sec2_intro = (
        f"藉由將「{city_prefix}{stage}{subject}{cohort_tag}學習能力檢測」題號，"
        f"對應之十二年國教課程綱要學習重點和答對率進行交叉對照，可更明確檢視學生表現較優異和有待加強之處，"
        f"相關整理如表5-1所示。"
    )
    p_desc2 = create_paragraph(sec2_intro, font_size_pt=12.0, line_spacing_pt=24.0, space_after_pt=6.0)
    ind2 = OxmlElement("w:ind")
    ind2.set(qn("w:firstLine"), "480")
    p_desc2.find(qn("w:pPr")).append(ind2)
    elements.append(p_desc2)
    
    # Group items by content domain/area
    area_dict: dict[str, list[Item]] = {}
    for item in items:
        area_name = item.area or "主要評量向度"
        area_dict.setdefault(area_name, []).append(item)
        
    sorted_area_names = sort_areas(list(area_dict.keys()), subject, grade_num, is_7n=is_7n)
    
    table_index = 1
    for area_name in sorted_area_names:
        area_items = area_dict[area_name]
        caption = f"表5-1-{table_index}　{area_name}學習重點及其答對率之交叉對照表"
        p_tbl_cap = create_paragraph(caption, font_size_pt=11.0, bold=True, space_after_pt=4.0)
        jc_tcap = OxmlElement("w:jc")
        jc_tcap.set(qn("w:val"), "center")
        p_tbl_cap.find(qn("w:pPr")).append(jc_tcap)
        elements.append(p_tbl_cap)
        
        # Build Table (3-column for Chinese, 4-column with merged header for Math/English/Science)
        tbl = build_area_cross_table(area_items, descriptions, is_chinese=is_chinese)
        elements.append(tbl)
        
        # Spacer
        elements.append(create_paragraph("", font_size_pt=12.0, space_after_pt=8.0))
        table_index += 1
        
    return elements


def create_standalone_chapter_5_docx(
    output_path: str | Path,
    subject: str,
    grade: int | str,
    items: list[Item],
    descriptions: dict[str, str],
    figure_path: str | Path | None = None,
    year: str = "115",
    is_7n: bool = False,
) -> Path:
    """Builds a standalone official Chapter 5 Word document."""
    from docx.shared import Pt
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Pt(64.8)
    sec.bottom_margin = Pt(50.4)
    sec.left_margin = Pt(56.7)
    sec.right_margin = Pt(49.6)
    
    elems = build_chapter_5(
        subject=subject,
        grade=grade,
        items=items,
        descriptions=descriptions,
        figure_path=figure_path,
        year=year,
        is_7n=is_7n,
    )
    for elem in elems:
        doc.element.body.append(elem)
        
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out
