"""SAAA Report Engine: Chapter 7 Appendix Module.
Generates:
  - 柒、附錄 (Chapter 7)
  - Full-page booklet PDF page rendering and injection.
  - Sub-document appendix embedding / standalone Chapter 7 document generation.
"""

from __future__ import annotations

import io
import os
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt, Emu
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from .ch1_2_admin import chinese_numeral, create_paragraph, set_run_font

DEFAULT_DPI = 181.4068
TARGET_WIDTH_EMU = 6210300   # 489.0 pt
TARGET_HEIGHT_EMU = 8785504  # 691.8 pt


def create_chapter_7_heading(text: str = "柒、附錄") -> OxmlElement:
    """Creates the official Chapter 7 heading paragraph with exact shading and 24pt line height."""
    xml = (
        f'<w:p {nsdecls("w")}>'
        '  <w:pPr>'
        '    <w:widowControl/>'
        '    <w:shd w:val="clear" w:color="auto" w:fill="DDEBF7"/>'
        '    <w:spacing w:line="480" w:lineRule="exact"/>'
        '    <w:ind w:left="482" w:hanging="482"/>'
        '    <w:jc w:val="center"/>'
        '    <w:textAlignment w:val="center"/>'
        '    <w:rPr>'
        '      <w:rFonts w:ascii="微軟正黑體" w:eastAsia="微軟正黑體" w:hAnsi="微軟正黑體"/>'
        '      <w:b/>'
        '      <w:spacing w:val="20"/>'
        '      <w:sz w:val="24"/>'
        '      <w:szCs w:val="24"/>'
        '    </w:rPr>'
        '  </w:pPr>'
        '  <w:r>'
        '    <w:rPr>'
        '      <w:rFonts w:ascii="微軟正黑體" w:eastAsia="微軟正黑體" w:hAnsi="微軟正黑體"/>'
        '      <w:b/>'
        '      <w:sz w:val="24"/>'
        '      <w:szCs w:val="24"/>'
        '    </w:rPr>'
        f'    <w:t>{text}</w:t>'
        '  </w:r>'
        '</w:p>'
    )
    return parse_xml(xml)


def render_pdf_to_image_paragraphs(
    pdf_path: str | Path,
    dpi: float = DEFAULT_DPI,
    width_emu: int = TARGET_WIDTH_EMU,
    height_emu: int = TARGET_HEIGHT_EMU,
) -> list[Any]:
    """Renders all pages of a booklet PDF into python-docx paragraph XML elements."""
    import pymupdf
    
    p_path = Path(pdf_path)
    if not p_path.exists():
        return []
        
    doc_pdf = pymupdf.open(str(p_path))
    elements = []
    
    scale = dpi / 72.0
    matrix = pymupdf.Matrix(scale, scale)
    
    temp_doc = Document()
    
    for page_idx in range(len(doc_pdf)):
        page = doc_pdf[page_idx]
        pix = page.get_pixmap(matrix=matrix, alpha=False)
        img_bytes = pix.tobytes("png")
        
        p = temp_doc.add_paragraph()
        p_pPr_xml = (
            f'<w:pPr {nsdecls("w")}>'
            '  <w:keepLines/>'
            + ('  <w:pageBreakBefore w:val="0"/>' if page_idx == 0 else '  <w:pageBreakBefore/>')
            + '  <w:spacing w:before="0" w:after="0"/>'
            '  <w:ind w:left="0" w:right="0"/>'
            '  <w:jc w:val="center"/>'
            '</w:pPr>'
        )
        p._p.append(parse_xml(p_pPr_xml))
        
        run = p.add_run()
        run.add_picture(io.BytesIO(img_bytes), width=Emu(width_emu), height=Emu(height_emu))
        elements.append(p._p)
        
    return elements


def build_chapter_7(
    subject: str,
    grade: int,
    appendix_files: list[str | Path] | None = None,
    booklet_pdf: str | Path | None = None,
    year: str = "115",
    official_style: bool = True,
) -> list[Any]:
    """Generates complete Chapter 7 (柒、附錄) elements.
    
    If booklet_pdf is provided, renders all pages directly into full-page appendix pictures.
    """
    elements = []
    stage = "國民小學" if grade <= 6 else "國民中學"
    g_cn = chinese_numeral(grade)
    sub_full = "自然科學領域" if subject == "自然" else f"{subject}領域" if "領域" not in subject else subject
    
    # 1. Heading: 柒、附錄
    if official_style:
        elements.append(create_chapter_7_heading("柒、附錄"))
    else:
        h7 = create_paragraph("柒、附錄", font_size_pt=16.0, bold=True, space_after_pt=12.0)
        elements.append(h7)
        app1_title = f"附錄一　{year}年度{stage}{sub_full}{g_cn}年級學生學習能力檢測試題本與參考答案"
        p_app1 = create_paragraph(app1_title, font_size_pt=14.0, bold=True, space_after_pt=6.0)
        elements.append(p_app1)
    
    inserted_any = False
    
    # Priority 1: Direct booklet PDF rendering
    if booklet_pdf and Path(booklet_pdf).exists():
        pdf_elems = render_pdf_to_image_paragraphs(booklet_pdf)
        if pdf_elems:
            elements.extend(pdf_elems)
            inserted_any = True
            
    # Priority 2: Appendix files list
    if not inserted_any and appendix_files:
        for f_path in appendix_files:
            p = Path(f_path)
            if not p.exists():
                continue
            if p.suffix.lower() == ".pdf":
                pdf_elems = render_pdf_to_image_paragraphs(p)
                if pdf_elems:
                    elements.extend(pdf_elems)
                    inserted_any = True
            elif p.suffix.lower() == ".docx":
                try:
                    ext_doc = Document(str(p))
                    for child in ext_doc.element.body.iterchildren():
                        if child.tag in (qn("w:p"), qn("w:tbl")):
                            elements.append(child)
                            inserted_any = True
                except Exception:
                    pass
            elif p.suffix.lower() in (".png", ".jpg", ".jpeg"):
                temp_doc = Document()
                img_p = temp_doc.add_paragraph()
                img_p.alignment = 1
                run = img_p.add_run()
                run.add_picture(str(p), width=Emu(TARGET_WIDTH_EMU), height=Emu(TARGET_HEIGHT_EMU))
                elements.append(img_p._p)
                inserted_any = True
                
    if not inserted_any:
        p_desc = create_paragraph(
            f"（本附錄收錄{year}年度{stage}{sub_full}{g_cn}年級之正式施測試題題本、參考答案及相關試題疑義回覆公告）",
            font_size_pt=11.0, line_spacing_pt=24.0, space_after_pt=12.0
        )
        elements.append(p_desc)
        
    return elements


def create_standalone_chapter_7_docx(
    pdf_path: str | Path,
    output_path: str | Path,
    dpi: float = DEFAULT_DPI,
) -> Path:
    """Generates a standalone Chapter 7 DOCX document adhering to official specifications."""
    doc = Document()
    
    # Official margins and page size:
    section = doc.sections[0]
    section.page_width = Pt(595.3)
    section.page_height = Pt(841.9)
    section.top_margin = Pt(64.8)
    section.bottom_margin = Pt(50.4)
    section.left_margin = Pt(56.7)
    section.right_margin = Pt(49.6)
    
    # Paragraph 0: Heading 柒、附錄
    p0 = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    p0._p.append(parse_xml(
        f'<w:pPr {nsdecls("w")}>'
        '  <w:widowControl/>'
        '  <w:shd w:val="clear" w:color="auto" w:fill="DDEBF7"/>'
        '  <w:spacing w:line="480" w:lineRule="exact"/>'
        '  <w:ind w:left="482" w:hanging="482"/>'
        '  <w:jc w:val="center"/>'
        '  <w:textAlignment w:val="center"/>'
        '  <w:rPr>'
        '    <w:rFonts w:ascii="微軟正黑體" w:eastAsia="微軟正黑體" w:hAnsi="微軟正黑體"/>'
        '    <w:b/>'
        '    <w:spacing w:val="20"/>'
        '    <w:sz w:val="24"/>'
        '    <w:szCs w:val="24"/>'
        '  </w:rPr>'
        '</w:pPr>'
    ))
    r0 = p0.add_run("柒、附錄")
    r0.bold = True
    r0.font.name = "微軟正黑體"
    r0.font.size = Pt(12)
    rPr = r0._r.get_or_add_rPr()
    rPr.append(parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="微軟正黑體" w:eastAsia="微軟正黑體" w:hAnsi="微軟正黑體"/>'))
    
    # Render PDF pages
    pdf_elems = render_pdf_to_image_paragraphs(pdf_path, dpi=dpi)
    for elem in pdf_elems:
        doc.element.body.append(elem)
        
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out
